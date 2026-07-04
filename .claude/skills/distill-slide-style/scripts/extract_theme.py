#!/usr/bin/env python3
"""Extract ground-truth theme/master/layout data from .pptx files.

Usage:
    extract_theme.py <file-or-dir> [<file-or-dir> ...] [--out theme_profile.json]

Unlike extract_pptx.py (which infers color/font conventions bottom-up from
per-run formatting actually used on slides), this script reads the
authoritative source: each file's slide master(s), the theme(s) they
reference (color scheme, font scheme including the CJK/Hans-script font
fallback, which python-pptx's font-scheme API doesn't expose directly), and
every named slide layout's placeholder geometry (the composition grid a
designer intended, independent of what any single slide happens to fill in).

Feed the resulting theme_profile.json into step 3 aggregation alongside the
per-sample style profiles — it lets the spec distinguish "the org's declared
brand palette/grid" from "what authors actually typed," and catch drift
between the two (e.g. an author picking PowerPoint's standard-palette swatch
instead of the linked theme color, which look near-identical but are not the
same thing).
"""
import argparse
import json
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A_NS}

COLOR_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
               "accent4", "accent5", "accent6", "hlink", "folHlink")


def theme_color_scheme(theme_xml: bytes) -> dict:
    root = etree.fromstring(theme_xml)
    scheme = root.find(".//a:clrScheme", NS)
    colors = {}
    if scheme is not None:
        for child in scheme:
            slot = etree.QName(child).localname
            if slot not in COLOR_SLOTS:
                continue
            srgb = child.find("a:srgbClr", NS)
            sysclr = child.find("a:sysClr", NS)
            colors[slot] = srgb.get("val") if srgb is not None else (
                sysclr.get("lastClr") if sysclr is not None else None
            )
    return colors


def theme_font_scheme(theme_xml: bytes) -> dict:
    root = etree.fromstring(theme_xml)
    result = {}
    for role, xpath in (("major", ".//a:fontScheme/a:majorFont"),
                        ("minor", ".//a:fontScheme/a:minorFont")):
        font_el = root.find(xpath, NS)
        if font_el is None:
            continue
        latin = font_el.find("a:latin", NS)
        # CJK typefaces live in per-script <a:font script="Hans"/> fallback
        # entries, not the generic <a:ea> slot, which is usually blank.
        hans = font_el.find("a:font[@script='Hans']", NS)
        result[role] = {
            "latin": latin.get("typeface") if latin is not None else None,
            "hans": hans.get("typeface") if hans is not None else None,
        }
    return result


def placeholder_geometry(layout, slide_w, slide_h) -> list:
    geometry = []
    for ph in layout.placeholders:
        pf = ph.placeholder_format
        entry = {"idx": pf.idx, "type": str(pf.type)}
        if ph.left is None:
            entry["position"] = "inherited"
        else:
            entry["position"] = {
                "left": round(ph.left / slide_w, 3),
                "top": round(ph.top / slide_h, 3),
                "width": round(ph.width / slide_w, 3),
                "height": round(ph.height / slide_h, 3),
            }
        geometry.append(entry)
    return geometry


def demo_slide_findings(prs) -> list:
    """Real slides in a template file often demonstrate the intended chart/
    table color convention directly (e.g. a labeled 'delete before real use'
    reference page) -- worth surfacing separately from ordinary content."""
    findings = []
    for i, slide in enumerate(prs.slides):
        texts = [
            s.text_frame.text.strip() for s in slide.shapes
            if getattr(s, "has_text_frame", False) and s.text_frame.text.strip()
        ]
        entry = {"slide_index": i, "layout": slide.slide_layout.name, "texts": texts[:5]}
        for shape in slide.shapes:
            if getattr(shape, "has_chart", False):
                series_colors = []
                for plot in shape.chart.plots:
                    for series in plot.series:
                        try:
                            series_colors.append(str(series.format.fill.fore_color.theme_color))
                        except Exception:
                            series_colors.append("explicit-or-inherited")
                entry["chart_series_theme_colors"] = series_colors
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                entry["has_table"] = True
        findings.append(entry)
    return findings


def extract_file(path: Path) -> dict:
    prs = Presentation(str(path))
    slide_w, slide_h = prs.slide_width, prs.slide_height

    themes = {}
    masters = []
    for master in prs.slide_masters:
        theme_part = None
        for rel in master.part.rels.values():
            if rel.reltype.endswith("/theme"):
                theme_part = rel.target_part
                break
        theme_name = theme_part.partname if theme_part else None
        if theme_part and str(theme_name) not in themes:
            themes[str(theme_name)] = {
                "colors": theme_color_scheme(theme_part.blob),
                "fonts": theme_font_scheme(theme_part.blob),
            }
        layouts = []
        for layout in master.slide_layouts:
            layouts.append({
                "name": layout.name,
                "placeholders": placeholder_geometry(layout, slide_w, slide_h),
            })
        masters.append({"theme_ref": str(theme_name), "layouts": layouts})

    return {
        "source": path.name,
        "slide_dimensions": {"width_in": round(slide_w / 914400, 2), "height_in": round(slide_h / 914400, 2)},
        "themes": themes,
        "masters": masters,
        "demo_slides": demo_slide_findings(prs),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument("--out", type=Path, default=Path("theme_profile.json"))
    args = parser.parse_args()

    files = []
    for p in args.paths:
        files.extend(sorted(p.rglob("*.pptx")) if p.is_dir() else [p])
    if not files:
        print("No .pptx files found", file=sys.stderr)
        sys.exit(1)

    result = {"files": [extract_file(f) for f in files]}
    args.out.write_text(json.dumps(result, ensure_ascii=False, indent=2))
    print(f"Wrote theme profile for {len(files)} file(s) to {args.out}")


if __name__ == "__main__":
    main()
